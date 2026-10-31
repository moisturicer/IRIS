"""
`POST /records/import_excel/` with a real .xlsx, through the HTTP boundary (IR-315).

The importer needs `pyexcel`, which requirements/base.txt never declared -- only
its plugins were -- so every upload answered "pyexcel is not installed" in
production. The tests that should have shown it were skipped behind a
`find_spec("pyexcel")` guard and were later deleted with the fixed-stage suites
(IR-260). Nothing here is skipped: a missing dependency must fail the build.
"""
from datetime import date
from io import BytesIO
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from core.enums import PipelineStatus, RecordTypeName
from .models import Author, Classification, PSCEDClassification, Record, RecordOwner, RecordType
from .tests import make_user

#: The importer's documented columns (`parse_excel_import`), not the template's
#: output, so a drift between the two cannot hide here.
HEADERS = [
    "Title", "Abstract", "Year Accomplished", "Year Completed", "Record Type",
    "Classification", "PSCED Classification", "Is IP", "For Commercialization",
    "Community Extension", "Authors",
]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def row(*cells):
    """The leading cells of one data row, blank to the last column."""
    return [*cells] + [None] * (len(HEADERS) - len(cells))


def workbook(*rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ImportTestCase(APITestCase):
    """Signed in as RDCO, uploading through the real endpoint."""

    def setUp(self):
        self.rdco = make_user("import_rdco@cit.edu", "RDCO")
        self.client.force_authenticate(self.rdco)

    def post(self, payload, name="legacy.xlsx"):
        return self.client.post(
            reverse("record-import-excel"),
            {"file": SimpleUploadedFile(name, payload, content_type=XLSX)},
            format="multipart",
        )

    def post_rows(self, *rows):
        return self.post(workbook(*rows))


class ExcelImportTests(ImportTestCase):
    def test_declared_external_origin_and_publication_year_survive_import(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Title", "Authors", "Origin", "Publication Year", "Year Accomplished"])
        # Verified 2026-10-30 against https://arxiv.org/abs/1706.03762.
        authors = ["Ashish Vaswani", "Noam Shazeer", "Niki Parmar", "Jakob Uszkoreit", "Llion Jones", "Aidan N. Gomez", "Lukasz Kaiser", "Illia Polosukhin"]
        wb.active.append(["Attention Is All You Need", ", ".join(authors), "imported_reference", 2017, 2020])
        buffer = BytesIO()
        wb.save(buffer)
        response = self.post(buffer.getvalue())
        self.assertEqual(response.status_code, 200)
        record = Record.objects.get(title="Attention Is All You Need")
        self.assertEqual((record.origin, record.publication_year, record.year_accomplished), ("imported_reference", 2017, 2020))
        self.assertEqual(list(record.authors.order_by("pk").values_list("name", flat=True)), authors)

    def test_invalid_origin_is_reported_and_never_reclassified(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Title", "Origin"])
        wb.active.append(["Bad provenance", "guessed"])
        buffer = BytesIO()
        wb.save(buffer)
        response = self.post(buffer.getvalue())
        self.assertEqual(response.data["created"], 0)
        self.assertIn("Origin", response.data["log"])

    def test_a_real_xlsx_upload_creates_a_published_record(self):
        response = self.post(workbook([
            "An Imported Legacy Record", "A legacy abstract.", 2021, 2022,
            RecordTypeName.PROJECT, None, None, "YES", "no", "TRUE",
            "Juan Dela Cruz, Maria Santos",
        ]))

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["created"], 1, response.data["log"])
        self.assertEqual(response.data["skipped"], 0, response.data["log"])

        record = Record.objects.get(title="An Imported Legacy Record")
        self.assertEqual(record.pipeline_status, PipelineStatus.PUBLISHED)
        self.assertEqual(record.record_type, RecordType.objects.get(name=RecordTypeName.PROJECT))
        self.assertEqual((record.year_accomplished, record.year_completed), (2021, 2022))
        self.assertEqual(
            (record.is_ip, record.for_commercialization, record.community_extension),
            (True, False, True),
        )
        self.assertTrue(record.owners.filter(user=self.rdco, is_primary=True).exists())
        self.assertEqual(
            sorted(Author.objects.filter(record=record).values_list("name", flat=True)),
            ["Juan Dela Cruz", "Maria Santos"],
        )

    def test_a_row_without_a_title_is_reported_by_row_and_the_rest_still_import(self):
        """
        Rows are independent: one bad row is named in the log by its sheet row
        number and skipped, and the rows around it are still created.
        """
        response = self.post(workbook(
            row("First Imported Record"),
            row(None, "An abstract with no title."),
            row("Third Imported Record"),
        ))

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual((response.data["created"], response.data["skipped"]), (2, 1))
        self.assertIn("Row 3: skipped — Title is required.", response.data["log"])
        self.assertEqual(
            set(Record.objects.filter(added_by=self.rdco).values_list("title", flat=True)),
            {"First Imported Record", "Third Imported Record"},
        )

    def test_a_file_that_is_not_a_spreadsheet_is_reported_and_creates_nothing(self):
        response = self.post(b"this is not a workbook", name="notes.xlsx")

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual((response.data["created"], response.data["skipped"]), (0, 1))
        self.assertIn("Could not open file", response.data["log"])
        self.assertFalse(Record.objects.filter(added_by=self.rdco).exists())


class ImportRowIsAllOrNothingTests(ImportTestCase):
    """
    A row is a Record, its primary owner and its authors, saved together or
    not at all (IR-637). The import publishes straight away, so a half-saved
    row was a public, unreviewed record the log said was never created.
    """

    def test_a_failed_author_write_leaves_nothing_of_its_row_and_spares_the_others(self):
        real_create = Author.objects.create

        def failing_for_one_row(**fields):
            if fields["name"] == "Broken Author":
                raise IntegrityError("author write refused")
            return real_create(**fields)

        payload = workbook(
            row("Row Before", None, None, None, None, None, None, None, None, None, "Ana Reyes"),
            row("The Failing Row", None, None, None, None, None, None, None, None, None,
                "Good Author, Broken Author"),
            row("Row After", None, None, None, None, None, None, None, None, None, "Ben Cruz"),
        )
        with patch.object(Author.objects, "create", side_effect=failing_for_one_row):
            response = self.post(payload)

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual((response.data["created"], response.data["skipped"]), (2, 1), response.data["log"])
        # Nothing of the failing row: no Record, so no owner and no author either.
        self.assertFalse(Record.objects.filter(title="The Failing Row").exists())
        self.assertFalse(Author.objects.filter(name__in=["Good Author", "Broken Author"]).exists())
        self.assertEqual(RecordOwner.objects.filter(user=self.rdco).count(), 2)
        self.assertIn("✗  Row 3: not saved — The Failing Row — author write refused", response.data["log"])
        self.assertEqual(
            set(Record.objects.filter(added_by=self.rdco).values_list("title", flat=True)),
            {"Row Before", "Row After"},
        )
        self.assertEqual(
            set(Author.objects.values_list("name", flat=True)), {"Ana Reyes", "Ben Cruz"},
        )

    def test_a_real_database_error_mid_row_is_rolled_back_and_the_import_continues(self):
        """
        No mock: Postgres refuses an author name longer than its column, which
        aborts the transaction. The row's savepoint is what lets the rows
        after it still be written.
        """
        too_long = "A" * 201

        response = self.post_rows(
            row("Before The Real Error"),
            row("Real Error Row", None, None, None, None, None, None, None, None, None, f"Fine Name, {too_long}"),
            row("After The Real Error"),
        )

        self.assertEqual((response.data["created"], response.data["skipped"]), (2, 1), response.data["log"])
        self.assertFalse(Record.objects.filter(title="Real Error Row").exists())
        self.assertFalse(Author.objects.filter(name="Fine Name").exists())
        line = next(l for l in response.data["log"].splitlines() if "Row 3:" in l)
        self.assertTrue(line.startswith("✗  Row 3: not saved — Real Error Row — "), line)
        self.assertEqual(
            set(Record.objects.filter(added_by=self.rdco).values_list("title", flat=True)),
            {"Before The Real Error", "After The Real Error"},
        )

    def test_a_failed_owner_write_leaves_no_published_record(self):
        with patch.object(RecordOwner.objects, "create", side_effect=IntegrityError("owner write refused")):
            response = self.post(workbook(row("Ownerless Record")))

        self.assertEqual((response.data["created"], response.data["skipped"]), (0, 1), response.data["log"])
        self.assertFalse(Record.objects.filter(title="Ownerless Record").exists())
        self.assertIn("✗  Row 2: not saved — Ownerless Record — owner write refused", response.data["log"])


class ImportRefusesValuesItCannotAcceptTests(ImportTestCase):
    """
    The import publishes without review, so a value the system cannot accept
    skips its row and says why, as a missing Title already does (IR-638).
    """

    def setUp(self):
        super().setUp()
        Classification.objects.get_or_create(name="Applied Sciences")
        PSCEDClassification.objects.get_or_create(name="Information and Communication Technologies")
        self.next_year = date.today().year + 1

    def imported(self):
        return set(Record.objects.filter(added_by=self.rdco).values_list("title", flat=True))

    def test_a_year_that_is_not_a_number_skips_the_row_and_names_it(self):
        response = self.post_rows(
            row("Typo Year", None, "20O4"),
            row("Underscore Year", None, "2_024"),
            row("Good Year", None, 2004),
        )

        self.assertEqual((response.data["created"], response.data["skipped"]), (1, 2), response.data["log"])
        self.assertIn('Row 3: skipped — Year Accomplished "2_024" is not a year', response.data["log"])
        self.assertIn(
            f'Row 2: skipped — Year Accomplished "20O4" is not a year from 1900 to {self.next_year}.',
            response.data["log"],
        )
        self.assertEqual(self.imported(), {"Good Year"})
        self.assertEqual(Record.objects.get(title="Good Year").year_accomplished, 2004)

    def test_a_year_outside_1900_to_next_year_skips_the_row(self):
        response = self.post_rows(
            row("Too Old", None, 1066),
            row("Too New", None, None, self.next_year + 1),
            row("Next Year Is Fine", None, self.next_year, self.next_year),
            row("1900 Is Fine", None, 1900),
        )

        self.assertEqual((response.data["created"], response.data["skipped"]), (2, 2), response.data["log"])
        self.assertIn(f'Row 2: skipped — Year Accomplished "1066" is not a year from 1900 to {self.next_year}.',
                      response.data["log"])
        self.assertIn(f'Row 3: skipped — Year Completed "{self.next_year + 1}" is not a year from 1900 to {self.next_year}.',
                      response.data["log"])
        self.assertEqual(self.imported(), {"Next Year Is Fine", "1900 Is Fine"})

    def test_a_completion_before_the_accomplishment_skips_the_row(self):
        response = self.post_rows(row("Backwards", None, 2022, 2020), row("Same Year", None, 2022, 2022))

        self.assertEqual((response.data["created"], response.data["skipped"]), (1, 1), response.data["log"])
        self.assertIn(
            'Row 2: skipped — Year Completed "2020" is earlier than Year Accomplished "2022".',
            response.data["log"],
        )
        self.assertEqual(self.imported(), {"Same Year"})

    def test_an_unknown_record_type_skips_the_row_and_lists_the_accepted_names(self):
        accepted = ", ".join(sorted(RecordType.objects.values_list("name", flat=True), key=str.lower))

        response = self.post_rows(row("Short Form Type", None, None, None, "Thesis/Research"))

        self.assertEqual((response.data["created"], response.data["skipped"]), (0, 1), response.data["log"])
        self.assertIn(
            f'Row 2: skipped — Record Type "Thesis/Research" is not one of: {accepted}.',
            response.data["log"],
        )
        self.assertFalse(Record.objects.filter(title="Short Form Type").exists())

    def test_names_match_ignoring_case_and_blank_names_keep_their_defaults(self):
        response = self.post_rows(
            row("Any Case", None, None, None, RecordTypeName.THESIS_RESEARCH.upper(),
                "applied sciences", "INFORMATION AND COMMUNICATION TECHNOLOGIES"),
            row("All Blank"),
        )

        self.assertEqual((response.data["created"], response.data["skipped"]), (2, 0), response.data["log"])
        any_case = Record.objects.get(title="Any Case")
        self.assertEqual(any_case.record_type.name, RecordTypeName.THESIS_RESEARCH)
        self.assertEqual(any_case.classification.name, "Applied Sciences")
        self.assertEqual(any_case.psced.name, "Information and Communication Technologies")
        blank = Record.objects.get(title="All Blank")
        # A blank Record Type keeps its documented default; blank names stay blank.
        self.assertEqual(blank.record_type.name, RecordTypeName.PROJECT)
        self.assertIsNone(blank.classification)
        self.assertIsNone(blank.psced)

    def test_an_unknown_classification_or_psced_skips_the_row_and_lists_the_accepted_names(self):
        response = self.post_rows(
            row("Unknown Classification", None, None, None, None, "Applied Science"),
            row("Unknown PSCED", None, None, None, None, None, "ICT"),
        )

        self.assertEqual((response.data["created"], response.data["skipped"]), (0, 2), response.data["log"])
        classifications = ", ".join(sorted(Classification.objects.values_list("name", flat=True), key=str.lower))
        psced_names = ", ".join(sorted(PSCEDClassification.objects.values_list("name", flat=True), key=str.lower))
        self.assertIn(
            f'Row 2: skipped — Classification "Applied Science" is not one of: {classifications}.',
            response.data["log"],
        )
        self.assertIn(
            f'Row 3: skipped — PSCED Classification "ICT" is not one of: {psced_names}.',
            response.data["log"],
        )

    def test_a_row_with_several_bad_cells_reports_every_one(self):
        response = self.post_rows(
            row(None, None, "20O4", 3024, "Thesis/Research", "Applied Science"),
            row("Still Imported"),
        )

        self.assertEqual((response.data["created"], response.data["skipped"]), (1, 1), response.data["log"])
        line = next(l for l in response.data["log"].splitlines() if l.startswith("Row 2:"))
        for problem in (
            "Title is required",
            'Year Accomplished "20O4"',
            'Year Completed "3024"',
            'Record Type "Thesis/Research"',
            'Classification "Applied Science"',
        ):
            self.assertIn(problem, line)
        self.assertEqual(self.imported(), {"Still Imported"})

    def test_the_template_imports_its_own_sample_row_and_not_its_notes_row(self):
        """
        The downloaded template's sample used a Record Type name the system
        does not have, and its notes row imported as a record titled
        "Required" (both seen in IR-315). The notes row's "e.g. 2024" is not a
        year, so it is now skipped with its reasons.
        """
        template = self.client.get(reverse("record-download-template"))
        self.assertEqual(template.status_code, status.HTTP_200_OK)

        response = self.post(template.content, name="template.xlsx")

        sample = Record.objects.get(title="Blockchain-Based Verification of Academic Credentials")
        self.assertEqual(sample.record_type.name, RecordTypeName.THESIS_RESEARCH, response.data["log"])
        self.assertFalse(Record.objects.filter(title="Required").exists())
        self.assertEqual((response.data["created"], response.data["skipped"]), (1, 1), response.data["log"])
        self.assertIn('Row 3: skipped — Year Accomplished "e.g. 2024"', response.data["log"])

    def test_origin_and_publication_year_are_reported_with_the_rest_of_the_row(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Title", "Year Accomplished", "Origin", "Publication Year"])
        wb.active.append(["Three Bad Cells", "20O4", "guessed", "1999.5"])
        wb.active.append(["Fine Reference", None, "imported_reference", 1687])
        buffer = BytesIO()
        wb.save(buffer)

        response = self.post(buffer.getvalue())

        self.assertEqual((response.data["created"], response.data["skipped"]), (1, 1), response.data["log"])
        line = next(l for l in response.data["log"].splitlines() if "Row 2:" in l)
        for column in ('Year Accomplished "20O4"', 'Origin "guessed"', 'Publication Year "1999.5"'):
            self.assertIn(column, line)
        self.assertEqual(Record.objects.get(title="Fine Reference").publication_year, 1687)

    def test_a_year_too_long_to_be_a_number_is_a_skipped_row_not_a_crash(self):
        response = self.post_rows(row("Endless Year", None, "9" * 400), row("Good Year", None, 2004))

        self.assertEqual(response.status_code, 200)
        self.assertEqual((response.data["created"], response.data["skipped"]), (1, 1), response.data["log"])
